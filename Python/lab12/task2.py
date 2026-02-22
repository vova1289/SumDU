"""
Програма для обробки даних з Excel-файлів
Варіант 12: Робота з Excel
"""

import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import os
import random  # <-- ПЕРЕМІСТИВ ІМПОРТ СЮДИ
from datetime import datetime

# Файли
input_file = 'sales_data.xlsx'
output_file = 'report.xlsx'

def create_test_data():
    """Створює тестові дані, якщо файл не знайдено"""
    print("\nСтворення тестових даних...")

    data = []
    dates = pd.date_range('2025-01-01', periods=50, freq='D')
    categories = ['Електроніка', 'Одяг', 'Продукти']
    products = ['Ноутбук', 'Телефон', 'Планшет', 'Книга', 'Футболка']

    for i in range(50):
        price = random.randint(100, 1000)
        qty = random.randint(1, 5)
        data.append({
            'Date': dates[i],
            'Product': random.choice(products),
            'Category': random.choice(categories),
            'Price': price,
            'Quantity': qty,
            'Total': price * qty
        })

    df = pd.DataFrame(data)
    print(f"Створено {len(df)} записів")
    return df

def load_data():
    """Завантаження даних"""
    if os.path.exists(input_file):
        try:
            df = pd.read_excel(input_file)
            print(f"Завантажено {len(df)} записів")
            return df
        except:
            print("Помилка читання файлу")
            return create_test_data()
    else:
        print("Файл не знайдено, створюю тестові дані")
        return create_test_data()

def show_info(df):
    """Базова інформація"""
    print("\n" + "="*50)
    print("ІНФОРМАЦІЯ ПРО ДАНІ")
    print("="*50)
    print(f"Всього записів: {len(df)}")
    print(f"Загальна сума: {df['Total'].sum():,.0f} грн")
    print(f"Середня сума: {df['Total'].mean():,.0f} грн")
    print(f"Максимальна сума: {df['Total'].max():,.0f} грн")
    print(f"Мінімальна сума: {df['Total'].min():,.0f} грн")
    print("\nПерші 5 рядків:")
    print(df.head())
    print("\nСтатистика:")
    print(df.describe())

def filter_data(df):
    """Фільтрація даних"""
    print("\n" + "="*50)
    print("ФІЛЬТРАЦІЯ")
    print("="*50)

    filtered = df.copy()

    # Фільтр за категорією
    cats = df['Category'].unique()
    print("Доступні категорії:", ', '.join(cats))
    cat = input("Введіть категорію (Enter - пропустити): ").strip()
    if cat and cat in cats:
        filtered = filtered[filtered['Category'] == cat]
        print(f"Вибрано категорію: {cat}")

    # Фільтр за мінімальною сумою
    min_sum = input("Мінімальна сума (Enter - пропустити): ").strip()
    if min_sum:
        try:
            min_sum = float(min_sum)
            filtered = filtered[filtered['Total'] >= min_sum]
            print(f"Сума >= {min_sum}")
        except:
            print("Помилка вводу")

    print(f"Знайдено {len(filtered)} записів")
    if len(filtered) > 0:
        print("\nРезультат (перші 10):")
        print(filtered.head(10))

    return filtered

def sort_data(df):
    """Сортування даних"""
    print("\n" + "="*50)
    print("СОРТУВАННЯ")
    print("="*50)

    if len(df) == 0:
        print("Немає даних")
        return df

    print("Поля для сортування:", ', '.join(df.columns))
    field = input("Введіть поле для сортування: ").strip()

    if field not in df.columns:
        print("Невірне поле")
        return df

    order = input("Порядок (1 - зростання, 2 - спадання): ").strip()
    asc = (order == '1')

    sorted_df = df.sort_values(by=field, ascending=asc)
    print(f"Відсортовано за {field}")
    print(sorted_df.head(10))

    return sorted_df

def aggregate_data(df):
    """Агрегація даних"""
    print("\n" + "="*50)
    print("АГРЕГАЦІЯ")
    print("="*50)

    results = {}

    # За категоріями
    if 'Category' in df.columns:
        cat_sum = df.groupby('Category')['Total'].sum().sort_values(ascending=False)
        cat_count = df.groupby('Category')['Total'].count()
        print("\nПРОДАЖІ ЗА КАТЕГОРІЯМИ:")
        print("-" * 40)
        for cat in cat_sum.index:
            print(f"{cat:15} {cat_sum[cat]:10,.0f} грн  ({cat_count[cat]} прод.)")
        results['category'] = cat_sum

    # За місяцями
    if 'Date' in df.columns:
        df_copy = df.copy()
        df_copy['Month'] = pd.to_datetime(df_copy['Date']).dt.month
        month_names = {1:'Січень', 2:'Лютий', 3:'Березень', 4:'Квітень',
                       5:'Травень', 6:'Червень', 7:'Липень', 8:'Серпень',
                       9:'Вересень', 10:'Жовтень', 11:'Листопад', 12:'Грудень'}
        df_copy['MonthName'] = df_copy['Month'].map(month_names)
        month_sum = df_copy.groupby('MonthName')['Total'].sum()
        month_count = df_copy.groupby('MonthName')['Total'].count()

        print("\nПРОДАЖІ ЗА МІСЯЦЯМИ:")
        print("-" * 40)
        for mon in month_sum.index:
            print(f"{mon:10} {month_sum[mon]:10,.0f} грн  ({month_count[mon]} прод.)")
        results['monthly'] = month_sum

    return results

def create_charts(df, results):
    """Створення графіків"""
    print("\n" + "="*50)
    print("ГРАФІКИ")
    print("="*50)

    try:
        # Графік 1: за категоріями
        if 'category' in results:
            plt.figure(figsize=(10, 5))
            results['category'].plot(kind='bar', color='skyblue', edgecolor='black')
            plt.title('Продажі за категоріями', fontsize=14)
            plt.ylabel('Сума (грн)', fontsize=12)
            plt.xlabel('Категорія', fontsize=12)
            plt.xticks(rotation=45)
            plt.grid(axis='y', alpha=0.3)
            plt.tight_layout()
            plt.savefig('category_chart.png')
            plt.show()
            print("Збережено category_chart.png")

        # Графік 2: за місяцями
        if 'monthly' in results:
            plt.figure(figsize=(10, 5))
            results['monthly'].plot(kind='line', marker='o', color='red', linewidth=2)
            plt.title('Динаміка продажів за місяцями', fontsize=14)
            plt.ylabel('Сума (грн)', fontsize=12)
            plt.xlabel('Місяць', fontsize=12)
            plt.grid(True, alpha=0.3)
            plt.tight_layout()
            plt.savefig('monthly_chart.png')
            plt.show()
            print("Збережено monthly_chart.png")

    except Exception as e:
        print("Помилка створення графіків:", e)

def save_to_excel(original, filtered, sorted_df, results):
    """Збереження в Excel"""
    print("\n" + "="*50)
    print("ЗБЕРЕЖЕННЯ")
    print("="*50)

    try:
        wb = Workbook()

        # Видаляємо стандартний аркуш
        wb.remove(wb.active)

        # Стилі
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")

        # Аркуш 1: оригінальні дані
        ws1 = wb.create_sheet("Оригінал")
        # Заголовки
        for col, col_name in enumerate(original.columns, 1):
            cell = ws1.cell(1, col, col_name)
            cell.font = header_font
            cell.fill = header_fill
        # Дані
        for r, row in enumerate(original.head(100).to_numpy(), 2):
            for c, val in enumerate(row, 1):
                ws1.cell(r, c, val)

        # Аркуш 2: результати
        ws2 = wb.create_sheet("Аналіз")

        row = 1
        if 'category' in results:
            ws2.cell(row, 1, "ПРОДАЖІ ЗА КАТЕГОРІЯМИ")
            ws2.cell(row, 1).font = Font(bold=True, size=14)
            row += 1
            ws2.cell(row, 1, "Категорія")
            ws2.cell(row, 2, "Сума (грн)")
            for cell in [ws2.cell(row, 1), ws2.cell(row, 2)]:
                cell.font = header_font
                cell.fill = header_fill
            row += 1

            for cat, val in results['category'].items():
                ws2.cell(row, 1, cat)
                ws2.cell(row, 2, val)
                row += 1
            row += 2

        if 'monthly' in results:
            ws2.cell(row, 1, "ПРОДАЖІ ЗА МІСЯЦЯМИ")
            ws2.cell(row, 1).font = Font(bold=True, size=14)
            row += 1
            ws2.cell(row, 1, "Місяць")
            ws2.cell(row, 2, "Сума (грн)")
            for cell in [ws2.cell(row, 1), ws2.cell(row, 2)]:
                cell.font = header_font
                cell.fill = header_fill
            row += 1

            for mon, val in results['monthly'].items():
                ws2.cell(row, 1, mon)
                ws2.cell(row, 2, val)
                row += 1

        wb.save(output_file)
        print(f"Збережено в {output_file}")

    except Exception as e:
        print("Помилка збереження:", e)

def menu():
    """Головне меню"""
    print("\n" + "="*50)
    print("МЕНЮ")
    print("="*50)
    print("1 - Показати інформацію")
    print("2 - Фільтрувати дані")
    print("3 - Сортувати дані")
    print("4 - Агрегація (підсумки)")
    print("5 - Побудувати графіки")
    print("6 - Зберегти в Excel")
    print("7 - Виконати все")
    print("0 - Вихід")

def main():
    print("="*50)
    print("АНАЛІЗ ПРОДАЖІВ")
    print("Варіант 12")
    print("="*50)

    # Завантаження даних
    df = load_data()
    filtered = None
    sorted_data = None
    results = {}

    while True:
        menu()
        choice = input("Ваш вибір: ").strip()

        if choice == '1':
            show_info(df)

        elif choice == '2':
            filtered = filter_data(df)

        elif choice == '3':
            if filtered is not None and len(filtered) > 0:
                sorted_data = sort_data(filtered)
            else:
                sorted_data = sort_data(df)

        elif choice == '4':
            if filtered is not None and len(filtered) > 0:
                results = aggregate_data(filtered)
            else:
                results = aggregate_data(df)

        elif choice == '5':
            if not results:
                results = aggregate_data(df)
            create_charts(df, results)

        elif choice == '6':
            if not results:
                results = aggregate_data(df)
            save_to_excel(df, filtered, sorted_data, results)

        elif choice == '7':
            print("\nВИКОНАННЯ ВСІХ ОПЕРАЦІЙ...")
            show_info(df)
            results = aggregate_data(df)
            create_charts(df, results)
            save_to_excel(df, df, df, results)
            print("\nВСЕ ВИКОНАНО!")

        elif choice == '0':
            print("\nДо побачення!")
            break

        else:
            print("Невірний вибір. Спробуйте ще раз.")

if __name__ == "__main__":
    main()
